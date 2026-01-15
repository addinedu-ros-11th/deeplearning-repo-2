import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

MD_FILE = "Presentation_Materials.md"
PPTX_FILE = "Presentation_Enhanced.pptx"

FONT_NAME = "Noto Sans CJK KR"
TITLE_COLOR = RGBColor(11, 31, 58)
ACCENT_COLOR = RGBColor(237, 125, 49)
TEXT_COLOR = RGBColor(34, 34, 34)
BACKGROUND_COLOR = RGBColor(250, 249, 247)


def strip_md_emphasis(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    return text


def parse_markdown(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    slides = []
    current_slide = None
    presentation_title = None
    subtitle_lines = []
    pending_image_caption = None

    img_re = re.compile(r'<img\s+[^>]*src="([^"]+)"[^>]*>')

    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            presentation_title = strip_md_emphasis(line[2:])
            continue
        if line.startswith("## "):
            if current_slide:
                slides.append(current_slide)
            current_slide = {
                "title": strip_md_emphasis(line[3:]),
                "blocks": [],
                "image": None,
                "caption": None,
            }
            continue
        if line.startswith("### "):
            if current_slide:
                current_slide["blocks"].append(
                    {"type": "subheading", "text": strip_md_emphasis(line[4:])}
                )
            continue
        if line.startswith("---"):
            continue
        if line.startswith("<br"):
            continue

        img_match = img_re.search(line)
        if img_match and current_slide:
            img_src = img_match.group(1)
            current_slide["image"] = img_src
            pending_image_caption = current_slide
            continue

        if line.startswith("<small>") and line.endswith("</small>") and pending_image_caption:
            caption = line.replace("<small>", "").replace("</small>", "")
            pending_image_caption["caption"] = strip_md_emphasis(caption)
            pending_image_caption = None
            continue

        if line.startswith("<"):
            continue

        if line.startswith("> "):
            if current_slide:
                current_slide["blocks"].append(
                    {"type": "quote", "text": strip_md_emphasis(line[2:])}
                )
            else:
                subtitle_lines.append(strip_md_emphasis(line[2:]))
            continue

        if re.match(r"^\d+\.\s+", line):
            if current_slide:
                content = strip_md_emphasis(re.sub(r"^\d+\.\s+", "", line))
                current_slide["blocks"].append({"type": "bullet", "text": content})
            continue

        if line.startswith("- "):
            if current_slide:
                current_slide["blocks"].append(
                    {"type": "bullet", "text": strip_md_emphasis(line[2:])}
                )
            continue

        if current_slide:
            current_slide["blocks"].append(
                {"type": "text", "text": strip_md_emphasis(line)}
            )
        else:
            subtitle_lines.append(strip_md_emphasis(line))

    if current_slide:
        slides.append(current_slide)

    return presentation_title, subtitle_lines, slides


def set_slide_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size, bold=False, color=TEXT_COLOR, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = FONT_NAME
    if align:
        p.alignment = align
    return tf


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = color
    return shape


def add_image_fit(slide, image_path, left, top, width, height):
    image = Image.open(image_path)
    img_w, img_h = image.size
    box_w = width
    box_h = height
    img_ratio = img_w / img_h
    box_ratio = box_w / box_h

    if img_ratio > box_ratio:
        target_w = box_w
        target_h = box_w / img_ratio
    else:
        target_h = box_h
        target_w = box_h * img_ratio

    x = left + (box_w - target_w) / 2
    y = top + (box_h - target_h) / 2
    slide.shapes.add_picture(image_path, x, y, width=target_w, height=target_h)


def create_presentation(title, subtitles, slides, output_path):
    prs = Presentation()

    # Title slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(title_slide, TITLE_COLOR)
    slide_w = prs.slide_width
    slide_h = prs.slide_height

    title_box = title_slide.shapes.add_textbox(
        Inches(0.9), Inches(2.1), slide_w - Inches(1.8), Inches(1.2)
    )
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title or "Presentation"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.font.name = FONT_NAME

    subtitle_text = "\n".join([s for s in subtitles if s])
    if subtitle_text:
        add_textbox(
            title_slide,
            Inches(0.9),
            Inches(3.3),
            slide_w - Inches(1.8),
            Inches(1.3),
            subtitle_text,
            18,
            color=RGBColor(230, 230, 230),
        )

    add_rect(title_slide, Inches(0), slide_h - Inches(0.3), slide_w, Inches(0.3), ACCENT_COLOR)

    # Content slides
    for slide_data in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_background(slide, BACKGROUND_COLOR)

        add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(0.12), ACCENT_COLOR)

        add_textbox(
            slide,
            Inches(0.6),
            Inches(0.35),
            prs.slide_width - Inches(1.2),
            Inches(0.6),
            slide_data["title"],
            28,
            bold=True,
            color=TITLE_COLOR,
        )

        content_left = Inches(0.7)
        content_top = Inches(1.2)
        content_width = prs.slide_width - Inches(1.4)
        content_height = Inches(5.6)

        image_path = slide_data.get("image")
        caption = slide_data.get("caption")
        if image_path:
            image_full = os.path.join(os.path.dirname(output_path), image_path)
            if os.path.exists(image_full):
                img_w = Inches(3.1)
                gap = Inches(0.3)
                text_w = content_width - img_w - gap
                text_box = slide.shapes.add_textbox(
                    content_left, content_top, text_w, content_height
                )
                text_tf = text_box.text_frame
                text_tf.clear()
                add_text_blocks(text_tf, slide_data["blocks"])

                img_left = content_left + text_w + gap
                img_top = content_top
                img_height = content_height - Inches(0.5)
                add_image_fit(slide, image_full, img_left, img_top, img_w, img_height)

                if caption:
                    add_textbox(
                        slide,
                        img_left,
                        content_top + img_height,
                        img_w,
                        Inches(0.4),
                        caption,
                        11,
                        color=TEXT_COLOR,
                    )
                continue

        text_box = slide.shapes.add_textbox(
            content_left, content_top, content_width, content_height
        )
        text_tf = text_box.text_frame
        text_tf.clear()
        add_text_blocks(text_tf, slide_data["blocks"])

    prs.save(output_path)
    print(f"Presentation saved to {output_path}")


def add_text_blocks(text_frame, blocks):
    for idx, block in enumerate(blocks):
        if idx == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        if block["type"] == "subheading":
            p.text = block["text"]
            p.font.bold = True
            p.font.size = Pt(20)
            p.space_before = Pt(8)
        elif block["type"] == "quote":
            p.text = f"\"{block['text']}\""
            p.font.italic = True
            p.font.size = Pt(18)
            p.font.color.rgb = TITLE_COLOR
        elif block["type"] == "bullet":
            p.text = f"- {block['text']}"
            p.font.size = Pt(18)
        else:
            p.text = block["text"]
            p.font.size = Pt(18)
        p.font.name = FONT_NAME
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(6)

if __name__ == "__main__":
    current_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(current_dir, MD_FILE)
    pptx_path = os.path.join(current_dir, PPTX_FILE)

    title, subtitle_lines, slides_data = parse_markdown(md_path)
    create_presentation(title, subtitle_lines, slides_data, pptx_path)
